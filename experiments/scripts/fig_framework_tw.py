#!/usr/bin/env python3
"""Framework figure of the trustworthiness computation (calib paper, Sec. III).

One display list (DL), two renderers:
  * SVG master  -> docs/calib/figs_svg/framework_tw.svg  (+ PDF/PNG via cairosvg)
  * PPTX copy   -> docs/calib/figs/framework_tw.pptx     (native, editable shapes)

Panel (a): per-epoch observation model (map inset + position/heading likelihood).
Panel (b): HMM trellis, three recursions, three conditioning sets (ab the main
           score, delta/alpha muted), and the audit / downstream block.

Run:  python3 experiments/scripts/fig_framework_tw.py
"""
import math, os, html

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT_SVG = os.path.join(ROOT, "docs/calib/figs_svg/framework_tw.svg")
OUT_PDF = os.path.join(ROOT, "docs/calib/figs/framework_tw.pdf")
OUT_PNG = os.path.join(ROOT, "docs/calib/figs/framework_tw.png")
OUT_PPTX = os.path.join(ROOT, "docs/calib/figs/framework_tw.pptx")

W, H = 1440, 684

# Okabe-Ito derived palette
BLUE, BLUE_L = "#0072B2", "#E8F1F8"
ORANGE, ORANGE_D, ORANGE_L = "#E69F00", "#9A6B00", "#FBF0DC"
GREEN, GREEN_D, GREEN_L = "#009E73", "#046C4E", "#E3F4EE"
DARK, MID, GRAY, LGRAY = "#333333", "#666666", "#999999", "#C9CFD1"
NET, CAND, MAPBG = "#A8B2B3", "#566573", "#F7F7F2"
NODE_E = "#6E8CA0"
FONT = "Arial, Helvetica, sans-serif"
CHAR_W = 0.56         # mean Arial advance / em (normal weight)
CHAR_W_B = 0.62


def est_width(parts, F):
    """Rough laid-out width of a rich-text run, shared by both renderers."""
    w = 0.0
    for t, st, _c in parts:
        toks = set(st.split())
        eff = F * (0.68 if ("sub" in toks or "sup" in toks) else 1.0)
        w += len(t) * eff * (CHAR_W_B if "b" in toks else CHAR_W)
    return w

DL = []  # display list, z-order = insertion order


def rrect(x, y, w, h, fill=None, stroke=None, sw=1.5, rx=8, dash=None, opacity=1.0):
    DL.append(dict(op="rrect", x=x, y=y, w=w, h=h, fill=fill, stroke=stroke,
                   sw=sw, rx=rx, dash=dash, opacity=opacity))


def circle(cx, cy, r, fill=None, stroke=None, sw=1.5, dash=None, opacity=1.0):
    DL.append(dict(op="circle", cx=cx, cy=cy, r=r, fill=fill, stroke=stroke,
                   sw=sw, dash=dash, opacity=opacity))


def poly(pts, stroke=DARK, sw=1.6, dash=None, arrow=False, opacity=1.0,
         fill=None, cap="round"):
    DL.append(dict(op="poly", pts=list(pts), stroke=stroke, sw=sw, dash=dash,
                   arrow=arrow, opacity=opacity, fill=fill, cap=cap))


def line(p0, p1, **kw):
    poly([p0, p1], **kw)


def arc_pts(cx, cy, r, a0, a1, n=24):
    """Screen angles in degrees (y down, 0 = +x, positive = clockwise-down)."""
    return [(cx + r * math.cos(math.radians(a0 + (a1 - a0) * k / n)),
             cy + r * math.sin(math.radians(a0 + (a1 - a0) * k / n)))
            for k in range(n + 1)]


def text(x, y, parts, size=15, color=DARK, anchor="start", rotate=None,
         opacity=1.0, spacing=None):
    """parts: str, or list of (txt, style[, color]); style tokens: b i sub sup u."""
    if isinstance(parts, str):
        parts = [(parts, "")]
    norm = []
    for p in parts:
        t, st = p[0], p[1]
        c = p[2] if len(p) > 2 else None
        norm.append((t, st, c))
    DL.append(dict(op="text", x=x, y=y, parts=norm, size=size, color=color,
                   anchor=anchor, rotate=rotate, opacity=opacity,
                   spacing=spacing))


def proj(z, p0, p1):
    """Orthogonal projection of point z onto segment p0-p1 (clamped)."""
    dx, dy = p1[0] - p0[0], p1[1] - p0[1]
    t = ((z[0] - p0[0]) * dx + (z[1] - p0[1]) * dy) / (dx * dx + dy * dy)
    t = max(0.0, min(1.0, t))
    return (p0[0] + t * dx, p0[1] + t * dy)


# ============================================================ panel (a)
def panel_a():
    text(24, 30, [("(a) Per-epoch observation model", "b")], size=21,
         color="#222222")
    text(24, 50, [("input:  GNSS fixes ", ""), ("z", "i"), ("i", "i sub"),
                  (" with per-epoch ", ""), ("σ", "i"), ("i", "i sub"),
                  (" (NMEA GST)  ·  road network ", ""), ("G", "i"),
                  (" + UBODT shortest-path table", "")],
         size=13.5, color="#777777")

    # ---- map inset
    rrect(36, 58, 510, 264, fill=MAPBG, stroke="#CCCCCC", sw=1.2, rx=10)
    s1 = ((60, 245), (520, 245))
    s2 = ((302, 63), (418, 158))
    s3 = ((205, 63), (205, 240))
    for a, b in (((55, 300), (170, 300)), ((455, 70), (455, 200)),
                 ((62, 78), (150, 58))):
        line(a, b, stroke=NET, sw=2.2)
    for a, b in (s1, s2, s3):
        line(a, b, stroke=CAND, sw=3.4)

    zi, zim1 = (300, 168), (150, 110)
    R_DISK, R_ERR = 104, 44
    circle(*zi, R_DISK, stroke="#7F8C8D", sw=1.6, dash="7,4")
    circle(*zi, R_ERR, stroke=BLUE, sw=1.8, dash="5,3")

    # previous fix + displacement
    circle(*zim1, 5.0, fill="#FFFFFF", stroke=BLUE, sw=2.0)
    text(128, 104, [("z", "i"), ("i−1", "i sub")], size=14, color=BLUE)
    line((155, 112), (294, 165), stroke=BLUE, sw=2.0, arrow=True)
    text(238, 134, [("v", "i"), ("i", "i sub")], size=14.5, color=BLUE,
         anchor="middle")

    # the fix, error circle label
    circle(*zi, 6.5, fill=BLUE, stroke="#FFFFFF", sw=1.2)
    text(311, 160, [("z", "i"), ("i", "i sub")], size=15.5, color=BLUE)
    text(348, 160, [("σ̄", "i"), ("i", "i sub")], size=14, color=BLUE)

    # projections + residual on s1
    x1 = proj(zi, *s1)   # (300, 245)
    x2 = proj(zi, *s2)
    x3 = proj(zi, *s3)
    for p in (x1, x2, x3):
        rrect(p[0] - 4, p[1] - 4, 8, 8, fill=BLUE, stroke="#FFFFFF", sw=1.0, rx=1.5)
    line((300, 241), (300, 176), stroke=BLUE, sw=1.5, dash="4,3")
    text(292, 214, [("d", "i"), ("i,1", "i sub")], size=13.5, color=BLUE,
         anchor="end")
    text(508, 238, [("s", "i"), ("1", "i sub")], size=15, color=CAND,
         anchor="end")
    text(424, 165, [("s", "i"), ("2", "i sub")], size=15, color=CAND)
    text(197, 236, [("s", "i"), ("3", "i sub")], size=15, color=CAND,
         anchor="end")

    # theta figure at x1: tangent e_1 + translated displacement
    v = (zi[0] - zim1[0], zi[1] - zim1[1])
    nv = math.hypot(*v)
    ux, uy = v[0] / nv * 58, v[1] / nv * 58
    line((300, 245), (364, 245), stroke=ORANGE, sw=2.0, arrow=True)
    line((300, 245), (300 + ux, 245 + uy), stroke=BLUE, sw=1.8, arrow=True)
    ang = math.degrees(math.atan2(uy, ux))
    poly(arc_pts(300, 245, 34, 0, ang), stroke=ORANGE, sw=1.6)
    text(371, 240, [("e", "i"), ("1", "i sub")], size=13.5, color=ORANGE_D)
    text(344, 272, [("θ", "i"), ("i,1", "i sub")], size=14,
         color=ORANGE_D)

    # search-disk label with leader
    line((227, 241), (162, 266), stroke="#7F8C8D", sw=1.1)
    text(56, 270, [("search disk  ", ""), ("r", "i"), (" = 300 m", "")],
         size=14, color="#5D6D6E")

    # map legend
    circle(72, 291, 6, stroke=BLUE, sw=1.6, dash="3,2")
    text(84, 296, [("σ̄", "i"), ("i", "i sub"),
                   (" error circle", "")], size=12.5, color="#5D6D6E")
    line((232, 291), (260, 291), stroke="#7F8C8D", sw=1.5, dash="6,3")
    text(266, 296, "search disk (r)", size=12.5, color="#5D6D6E")
    line((60, 309), (88, 309), stroke=CAND, sw=3.0)
    text(94, 313, "candidate segment", size=12.5, color="#5D6D6E")
    line((232, 309), (260, 309), stroke=NET, sw=2.2)
    text(266, 313, "other network", size=12.5, color="#5D6D6E")

    # ---- position likelihood box
    rrect(36, 336, 530, 90, fill=BLUE_L, stroke=BLUE, sw=1.8, rx=9)
    text(52, 360, [("position likelihood", "b"),
                   ("  — isotropic Gaussian at σ̄", "b"),
                   ("i", "b i sub")], size=15.5, color=BLUE)
    text(52, 388, [("log p", "i"), ("pos", "i sub"), ("(z", "i"),
                   ("i", "i sub"), (" | x", "i"), ("i,j", "i sub"), (") = −‖d", "i"),
                   ("i,j", "i sub"), ("‖", "i"), ("2", "sup"),
                   (" / (2 σ̄", "i"), ("i", "i sub"), ("", "i"),
                   ("2", "sup"), (")  + const", "i")], size=15, color="#1A1A1A")
    text(52, 414, [("σ̄", "i"), ("i", "i sub"), ("2", "sup"),
                   (" = σ", "i"), ("i", "i sub"), ("2", "sup"),
                   (" + σ", "i"), ("map", "i sub"), ("2", "sup"),
                   (",   σ", "i"), ("map", "i sub"),
                   (" = √(10² + 3.5²) = 10.6 m", "i"),
                   ("   (map + lane, in variance)", "")],
         size=13.8, color="#33536B")

    # ---- heading likelihood box
    rrect(36, 448, 530, 102, fill=ORANGE_L, stroke=ORANGE, sw=1.8, rx=9)
    text(52, 472, [("heading likelihood", "b"),
                   ("  — one-sided von Mises penalty", "b")],
         size=15.5, color=ORANGE_D)
    text(52, 500, [("Λ", "i"), ("i,j", "i sub"), (" = κ", "i"),
                   ("i", "i sub"), (" cos θ", "i"), ("i,j", "i sub"),
                   ("  if cos θ", "i"), ("i,j", "i sub"),
                   (" < 0, else 0", "i"),
                   ("        κ", "i"), ("i", "i sub"), (" = ‖v", "i"),
                   ("i", "i sub"), ("‖", "i"), ("2", "sup"),
                   (" / (c σ", "i"), ("i", "i sub"), ("2", "sup"),
                   (" + b ‖v", "i"), ("i", "i sub"), ("‖", "i"),
                   ("2", "sup"), (")", "i")], size=14.5, color="#1A1A1A")
    text(52, 524, [("c", "i"), (" = 0.227,  ", ""), ("b", "i"),
                   (" = 0.0084 rad² — fitted · heading sees σ", ""),
                   ("i", "i sub"), (", not σ̄", ""), ("i", "i sub")],
         size=13, color="#7A5B12")
    text(52, 543, [("only reversed candidates discounted · e", ""),
                   ("j", "i sub"),
                   (" = segment tangent, not chord", "")],
         size=12.3, color="#8A6D1F")

    # ---- merge + layer normalisation (P box, Sigma, H box stacked)
    circle(300, 437, 10, fill=DARK)
    text(300, 441.5, [("Σ", "")], size=12, color="#FFFFFF", anchor="middle")
    text(318, 442, [("log p = log p", "i"), ("pos", "i sub"),
                    (" + Λ", "i"), ("i,j", "i sub")], size=13.5, color=DARK)
    line((300, 550), (300, 592), stroke=DARK, sw=1.8, arrow=True)
    rrect(118, 596, 364, 54, fill="#F0F0EE", stroke="#8A8A8A", sw=1.6, rx=9)
    text(300, 618, [("layer normalisation", "b")], size=14.5, color="#3A3A3A",
         anchor="middle")
    text(300, 639, [("ep", "i"), ("t", "i sub"), ("(j)", "i"),
                    (" = softmax over layer candidates", "")],
         size=13.2, color="#555555", anchor="middle")
    text(36, 670, [("empty disk → epoch unmatched & reported · invalid covariance → 5 m fallback", "")],
         size=11.8, color=GRAY)

# ============================================================ panel (b)
def panel_b():
    text(586, 30, [("(b) HMM inference — one run, one path, three conditioning sets", "b")],
         size=21, color="#222222")

    # ---- ep feed elbow from box E into the first layer
    poly([(482, 623), (608, 623), (608, 240), (686, 240)], stroke="#4A6A80",
         sw=2.0, arrow=True)
    text(598, 430, [("ep", "i"), ("t", "i sub"), ("(j) into every layer", "")],
         size=13.5, color="#4A6A80", anchor="middle", rotate=-90)

    # ---- forward band
    line((648, 78), (1112, 78), stroke=BLUE, sw=2.4, arrow=True)
    text(880, 64, [("forward ", ""), ("α", "i"), ("t", "i sub"),
                   (" — sum over all paths · causal, conditioned on ", ""),
                   ("z", "i"), ("1:t", "i sub")],
         size=14.5, color=BLUE, anchor="middle")

    # ---- trellis
    cols = {
        "tm1": (700, [145, 225, 305, 372]),
        "t":   (880, [125, 205, 290, 360]),
        "tp1": (1060, [150, 235, 310, 378]),
    }
    vit = [(700, 225), (880, 205), (1060, 235)]
    thin = [("tm1", 0, "t", 0), ("tm1", 0, "t", 1), ("tm1", 1, "t", 2),
            ("tm1", 2, "t", 2), ("tm1", 2, "t", 3), ("tm1", 3, "t", 3),
            ("tm1", 3, "t", 2),
            ("t", 0, "tp1", 0), ("t", 1, "tp1", 0), ("t", 2, "tp1", 1),
            ("t", 2, "tp1", 2), ("t", 3, "tp1", 2), ("t", 3, "tp1", 3)]
    for c1, i1, c2, i2 in thin:
        x1, ys1 = cols[c1]
        x2, ys2 = cols[c2]
        line((x1, ys1[i1]), (x2, ys2[i2]), stroke="#C5CCCE", sw=1.1)
    line(vit[0], vit[1], stroke=DARK, sw=3.6)
    line(vit[1], vit[2], stroke=DARK, sw=3.6)
    for cname, (x, ys) in cols.items():
        for y in ys:
            if (x, y) in vit:
                circle(x, y, 9.5, fill=DARK, stroke="#FFFFFF", sw=1.2)
            else:
                circle(x, y, 7.5, fill="#FFFFFF", stroke=NODE_E, sw=1.8)

    # i* highlight
    circle(880, 205, 15, stroke=GREEN, sw=2.6)
    text(900, 191, [("i*", "i b"), ("t", "i b sub")], size=15.5, color=GREEN_D)
    # viterbi label under the thick edge
    line((796, 232), (790, 220), stroke=GRAY, sw=1.1)
    text(800, 240, [("Viterbi ", ""), ("δ", "i"), ("t", "i sub"),
                    (" (max) selects ", ""), ("x*", "i"),
                    ("1:T", "i sub")], size=13, color=DARK)

    # epoch labels + fade
    for x, lab in ((700, "t−1"), (880, "t"), (1060, "t+1")):
        text(x, 398, [(lab, "i")], size=15, color=MID, anchor="middle")
    text(1122, 252, "…", size=26, color="#AAAAAA", anchor="middle")

    # ---- backward band
    line((1112, 406), (648, 406), stroke=GREEN, sw=2.4, dash="7,4", arrow=True)
    text(880, 428, [("backward ", ""), ("β", "i"), ("t", "i sub"),
                    (" — sum over future paths · anti-causal, ", ""),
                    ("z", "i"), ("t+1:T", "i sub")],
         size=14.5, color=GREEN, anchor="middle")

    # ---- transition / prior annotation
    text(624, 450, [("tp(a→b) = min(1, d", "i"), ("gnss", "i sub"),
                    (" / d", "i"), ("road", "i sub"), ("(a,b))", "i"),
                    (" — row-normalised · π = 1/K on the first layer", "")],
         size=14, color="#555555")

    # ---- outputs: muted delta/alpha, main alphabeta box
    rrect(1150, 118, 280, 80, fill="#FBFBFA", stroke="#BBBBBB", sw=1.2,
          rx=8, dash="5,3")
    text(1164, 142, [("tw", "i"), ("t", "i sub"), ("δ", "sup"),
                     (" — Viterbi score: max over paths", "")],
         size=12.8, color="#909090")
    text(1164, 164, [("tw", "i"), ("t", "i sub"), ("α", "sup"),
                     (" — filtering: causal, ", ""), ("z", "i"),
                     ("1:t", "i sub")], size=12.8, color="#909090")
    text(1164, 186, "same run · same path · same candidate",
         size=11.8, color="#AAAAAA")

    rrect(1150, 210, 280, 168, fill=GREEN_L, stroke=GREEN, sw=2.6, rx=10)
    text(1166, 238, [("smoothed posterior", "b")], size=16.5, color=GREEN_D)
    text(1166, 258, "the score this paper reports and audits", size=12.8,
         color="#4B8A72")
    text(1298, 290, [("tw", "i"), ("t", "i sub"), ("αβ", "sup"),
                     ("  =  α", "i"), ("t", "i sub"), ("(i*) · β", "i"),
                     ("t", "i sub"), ("(i*)", "i")],
         size=15.5, color="#1A1A1A", anchor="middle")
    line((1188, 300), (1408, 300), stroke="#1A1A1A", sw=1.6)
    text(1298, 322, [("Σ", "i"), ("j", "i sub"), (" α", "i"),
                     ("t", "i sub"), ("(j) · β", "i"), ("t", "i sub"),
                     ("(j)", "i")], size=15.5, color="#1A1A1A",
         anchor="middle")
    text(1166, 346, [("conditioned on the full trajectory ", ""), ("z", "i"),
                     ("1:T", "i sub")], size=12.8, color="#333333")
    text(1166, 368, [("offline · ECE 0.007 · BSS +0.27", "b")], size=12.8,
         color=GREEN_D)

    # connector i* -> main box
    poly([(895, 205), (1140, 205), (1140, 240), (1146, 240)], stroke=GREEN,
         sw=2.0, dash="6,4", arrow=True)


# ============================================================ audit block
def audit():
    # reliability thumbnail
    line((660, 490), (660, 600), stroke="#777777", sw=1.4)
    line((660, 600), (790, 600), stroke="#777777", sw=1.4)
    line((660, 600), (790, 490), stroke="#999999", sw=1.3, dash="4,3")
    us = [0, .125, .25, .375, .5, .625, .75, .875, 1]
    fs = [0, .10, .26, .40, .53, .68, .85, .95, 1]
    poly([(660 + 130 * u, 600 - 110 * f) for u, f in zip(us, fs)],
         stroke=GREEN, sw=2.6)
    text(794, 614, [("tw", "i"), ("t", "i sub")], size=12, color="#777777")
    text(652, 545, "freq. correct", size=11.5, color="#777777",
         anchor="middle", rotate=-90)
    text(725, 634, "reliability diagram (schematic)", size=12.5,
         color="#777777", anchor="middle")

    # audit box
    rrect(830, 485, 260, 120, fill="#F5F5F3", stroke="#888888", sw=1.6, rx=9)
    text(846, 509, [("calibration audit", "b")], size=15, color="#444444")
    text(846, 531, [("tw", "i"), ("t", "i sub"),
                    (" vs ground truth ", ""), ("c", "i"), ("t", "i sub"),
                    (":", "")], size=12.8, color="#555555")
    text(846, 553, [("ECE / MCE", "i"), (" — bin agreement", "")],
         size=13, color="#555555")
    text(846, 573, [("Brier / BSS", "i"), (" — proper score", "")],
         size=13, color="#555555")
    text(846, 593, [("ROC AUC / separation", "i")], size=13, color="#555555")

    # downstream box
    rrect(1140, 485, 274, 120, fill=GREEN_L, stroke=GREEN, sw=1.8, rx=9)
    text(1156, 509, [("downstream: calibrated weight", "b")], size=14.5,
         color=GREEN_D)
    text(1156, 533, [("M̂ = Σ", "i"), ("t", "i sub"), (" B", "i"),
                     ("t", "i sub"), (" (1 − tw", "i"), ("t", "i sub"),
                     (")", "i")], size=14, color="#1A1A1A")
    text(1156, 556, [("E[M̂] = E[M]", "i"), (" when ", ""),
                     ("E[c", "i"), ("t", "i sub"), (" | tw", "i"),
                     ("t", "i sub"), ("] = tw", "i"), ("t", "i sub")],
         size=12.5, color="#333333")
    text(1156, 580, "bias bounded by the deficit (Prop. 1)", size=11.8,
         color="#4B8A72")

    line((794, 545), (826, 545), stroke=MID, sw=2.0, arrow=True)
    line((1094, 545), (1136, 545), stroke=MID, sw=2.0, arrow=True)
    # tw output -> audit
    poly([(1290, 378), (1290, 466), (960, 466), (960, 483)], stroke=GREEN,
         sw=2.0, arrow=True)
    text(1298, 430, [("tw", "i b"), ("t", "i b sub")], size=14, color=GREEN_D)

# ============================================================ SVG renderer
def render_svg():
    arrow_cols = sorted({d["stroke"] for d in DL
                         if d["op"] == "poly" and d.get("arrow")})
    defs = "".join(
        f'<marker id="arw{c[1:]}" markerWidth="7" markerHeight="7" refX="6.2" '
        f'refY="3.5" orient="auto" markerUnits="strokeWidth">'
        f'<path d="M0,0 L7,3.5 L0,7 z" fill="{c}"/></marker>'
        for c in arrow_cols)
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" '
           f'viewBox="0 0 {W} {H}" font-family="{FONT}">',
           f'<defs>{defs}</defs>',
           f'<rect x="0" y="0" width="{W}" height="{H}" fill="#FFFFFF"/>']
    for d in DL:
        op = d["op"]
        if op == "rrect":
            a = (f'<rect x="{d["x"]}" y="{d["y"]}" width="{d["w"]}" '
                 f'height="{d["h"]}" rx="{d["rx"]}" '
                 f'fill="{d["fill"] or "none"}"')
            if d["stroke"]:
                a += (f' stroke="{d["stroke"]}" stroke-width="{d["sw"]}"')
                if d.get("dash"):
                    a += f' stroke-dasharray="{d["dash"]}"'
            if d["opacity"] != 1.0:
                a += f' opacity="{d["opacity"]}"'
            out.append(a + "/>")
        elif op == "circle":
            a = (f'<circle cx="{d["cx"]}" cy="{d["cy"]}" r="{d["r"]}" '
                 f'fill="{d["fill"] or "none"}"')
            if d["stroke"]:
                a += f' stroke="{d["stroke"]}" stroke-width="{d["sw"]}"'
                if d.get("dash"):
                    a += f' stroke-dasharray="{d["dash"]}"'
            if d["opacity"] != 1.0:
                a += f' opacity="{d["opacity"]}"'
            out.append(a + "/>")
        elif op == "poly":
            pts = " ".join(f"{x:.2f},{y:.2f}" for x, y in d["pts"])
            a = (f'<polyline points="{pts}" fill="{d["fill"] or "none"}" '
                 f'stroke="{d["stroke"]}" stroke-width="{d["sw"]}" '
                 f'stroke-linecap="round" stroke-linejoin="round"')
            if d.get("dash"):
                a += f' stroke-dasharray="{d["dash"]}"'
            if d["opacity"] != 1.0:
                a += f' opacity="{d["opacity"]}"'
            if d.get("arrow"):
                a += f' marker-end="url(#arw{d["stroke"][1:]})"'
            out.append(a + "/>")
        elif op == "text":
            F = d["size"]
            # cairosvg centres EVERY tspan under text-anchor="middle", so
            # non-start anchors are emulated by shifting x to the left edge.
            x = d["x"]
            if d["anchor"] != "start":
                tw_ = est_width(d["parts"], F)
                x = x - tw_ / 2 if d["anchor"] == "middle" else x - tw_
            a = (f'<text x="{x:.2f}" y="{d["y"]}" font-size="{F}" '
                 f'fill="{d["color"]}" text-anchor="start" '
                 f'xml:space="preserve"')
            if d["opacity"] != 1.0:
                a += f' opacity="{d["opacity"]}"'
            if d.get("rotate"):
                a += (f' transform="rotate({d["rotate"]} '
                      f'{d["x"]} {d["y"]})"')
            a += ">"
            dy_cur = 0.0
            for t, st, c in d["parts"]:
                toks = set(st.split())
                sub, sup = "sub" in toks, "sup" in toks
                sz = F * 0.68 if (sub or sup) else F
                tgt = 0.30 * F if sub else (-0.38 * F if sup else 0.0)
                dy = tgt - dy_cur
                dy_cur = tgt
                ta = ""
                if abs(sz - F) > 1e-9:
                    ta += f' font-size="{sz:.2f}"'
                if abs(dy) > 1e-9:
                    ta += f' dy="{dy:.2f}"'
                if c:
                    ta += f' fill="{c}"'
                if "b" in toks:
                    ta += ' font-weight="bold"'
                if "i" in toks:
                    ta += ' font-style="italic"'
                a += f"<tspan{ta}>{html.escape(t, quote=False)}</tspan>"
            out.append(a + "</text>")
    out.append("</svg>")
    return "\n".join(out)


# ============================================================ PPTX renderer
EMU_PX = 9525
TOP_K = 1.02          # baseline ≈ textbox top + TOP_K * fontsize (tune once)


def build_pptx(path):
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.dml import MSO_LINE_DASH_STYLE as MSO_LINE
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn

    def emu(v):
        return Emu(int(round(v * EMU_PX)))

    def rgb(c):
        return RGBColor.from_string(c.lstrip("#"))

    def style(shp, d):
        if d.get("fill"):
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(d["fill"])
        else:
            shp.fill.background()
        if d.get("stroke"):
            shp.line.color.rgb = rgb(d["stroke"])
            shp.line.width = Pt(d["sw"] * 0.75)
            if d.get("dash"):
                shp.line.dash_style = (MSO_LINE.ROUND_DOT if d["dash"] == "3,2"
                                       else MSO_LINE.DASH)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False

    def add_tail(shp):
        ln = shp.line._get_or_add_ln()
        te = ln.makeelement(qn("a:tailEnd"),
                            {"type": "triangle", "w": "med", "len": "med"})
        ln.append(te)

    prs = Presentation()
    prs.slide_width = emu(W)
    prs.slide_height = emu(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shapes = slide.shapes

    bg = shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, emu(W), emu(H))
    bg.fill.solid(); bg.fill.fore_color.rgb = rgb("#FFFFFF")
    bg.line.fill.background(); bg.shadow.inherit = False

    for d in DL:
        op = d["op"]
        if op == "rrect":
            shp = shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE if d["rx"] > 0
                else MSO_SHAPE.RECTANGLE,
                emu(d["x"]), emu(d["y"]), emu(d["w"]), emu(d["h"]))
            if d["rx"] > 0:
                try:
                    shp.adjustments[0] = min(0.5, d["rx"] /
                                             min(d["w"], d["h"]))
                except Exception:
                    pass
            style(shp, d)
        elif op == "circle":
            shp = shapes.add_shape(MSO_SHAPE.OVAL, emu(d["cx"] - d["r"]),
                                   emu(d["cy"] - d["r"]), emu(2 * d["r"]),
                                   emu(2 * d["r"]))
            style(shp, d)
        elif op == "poly":
            pts = d["pts"]
            for k in range(len(pts) - 1):
                conn = shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT, emu(pts[k][0]), emu(pts[k][1]),
                    emu(pts[k + 1][0]), emu(pts[k + 1][1]))
                conn.line.color.rgb = rgb(d["stroke"])
                conn.line.width = Pt(d["sw"] * 0.75)
                if d.get("dash"):
                    conn.line.dash_style = (MSO_LINE.ROUND_DOT if d["dash"] == "3,2"
                                            else MSO_LINE.DASH)
                conn.shadow.inherit = False
                if d.get("arrow") and k == len(pts) - 2:
                    add_tail(conn)
        elif op == "text":
            F = d["size"]
            tw = est_width(d["parts"], F) + 2
            th = F * 1.6
            if d.get("rotate"):
                left, top = d["x"] - tw / 2, d["y"] - th / 2
            else:
                left = {"start": d["x"], "middle": d["x"] - tw / 2,
                        "end": d["x"] - tw}[d["anchor"]]
                top = d["y"] - F * TOP_K
            tb = shapes.add_textbox(emu(left), emu(top), emu(tw), emu(th))
            if d.get("rotate"):
                tb.rotation = d["rotate"]
            tf = tb.text_frame
            tf.word_wrap = False
            tf.margin_left = tf.margin_right = 0
            tf.margin_top = tf.margin_bottom = 0
            tb.shadow.inherit = False
            p = tf.paragraphs[0]
            for t, st, c in d["parts"]:
                toks = set(st.split())
                sub, sup = "sub" in toks, "sup" in toks
                r = p.add_run()
                r.text = t
                f = r.font
                f.name = "Arial"
                f.size = Pt(F * (0.68 if (sub or sup) else 1.0) * 0.75)
                f.bold = "b" in toks
                f.italic = "i" in toks
                f.color.rgb = rgb(c or d["color"])
                if sub or sup:
                    r._r.get_or_add_rPr().set(
                        "baseline", "-25000" if sub else "30000")
    prs.save(path)


# ============================================================ main
def main():
    panel_a()
    panel_b()
    audit()
    for p in (OUT_SVG, OUT_PDF, OUT_PNG, OUT_PPTX):
        os.makedirs(os.path.dirname(p), exist_ok=True)
    svg = render_svg()
    with open(OUT_SVG, "w", encoding="utf-8") as fh:
        fh.write(svg)
    import cairosvg
    cairosvg.svg2pdf(url=OUT_SVG, write_to=OUT_PDF)
    cairosvg.svg2png(url=OUT_SVG, write_to=OUT_PNG, scale=2.0)
    build_pptx(OUT_PPTX)
    print(f"DL items: {len(DL)}")
    for p in (OUT_SVG, OUT_PDF, OUT_PNG, OUT_PPTX):
        print("wrote", p, os.path.getsize(p), "bytes")


if __name__ == "__main__":
    main()
